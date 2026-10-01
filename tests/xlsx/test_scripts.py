"""Tests for the xlsx skill's recalculation helper.

The shared `office/` tree is covered by the contract. What is specific here is
`recalc`, which drives LibreOffice to recompute formulas -- and the safety
check in front of it.

That check is the interesting part. Recalculating a workbook whose formulas
point at *other* workbooks, when those workbooks are unavailable, replaces
cached values with errors: the recalculation destroys the only copy of the
numbers. `external_links_at_risk` is what stops that, so the tests build real
workbooks with openpyxl and assert on exactly which cells it names.
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from pathlib import Path

import pytest

import skill_contract

SKILL_ROOT = Path(__file__).resolve().parents[2] / "skills" / "xlsx"
SCRIPTS = SKILL_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

openpyxl = pytest.importorskip("openpyxl", reason="xlsx scripts need openpyxl")

import recalc  # noqa: E402

OfficeTests = skill_contract.office.office_test_case(SKILL_ROOT)
CliHelpTests = skill_contract.cli.help_test_case(SKILL_ROOT)


class CommandLineTests(unittest.TestCase):

    def _run(self, *args: str):
        import os
        import subprocess

        return subprocess.run(
            [sys.executable, str(SCRIPTS / "recalc.py"), *args],
            capture_output=True,
            text=True,
            timeout=120,
            env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"},
        )

    def test_no_arguments_prints_usage(self) -> None:
        result = self._run()
        self.assertIn("usage:", result.stderr)
        self.assertEqual(result.returncode, 2)
        self.assertNotIn("Traceback", result.stderr)

    def test_the_usage_text_documents_the_force_flag(self) -> None:
        # --force is the escape hatch past the external-link guard; it has to
        # be discoverable or people will not know the guard can be overridden.
        self.assertIn("--force", self._run("--help").stdout)

    def test_invalid_timeout_reports_usage_without_traceback(self) -> None:
        result = self._run("book.xlsx", "nonsense")
        self.assertEqual(result.returncode, 2)
        self.assertNotIn("Traceback", result.stderr)


class ExternalReferencePatternTests(unittest.TestCase):
    """`EXTERNAL_REF_RE` decides which formulas count as reaching outside."""

    def test_bracketed_workbook_references_match(self) -> None:
        for formula in (
            "=[1]Sheet1!A1",
            "='[1]Some Sheet'!A1",
            "=SUM([2]Data!A1:A9)",
        ):
            with self.subTest(formula=formula):
                self.assertTrue(recalc.EXTERNAL_REF_RE.search(formula), formula)

    def test_ordinary_local_formulas_do_not_match(self) -> None:
        for formula in (
            "=A1+B2",
            "=SUM(Sheet2!A1:A9)",
            "=IF(A1>0,1,0)",
            '=CONCATENATE("[1]",A1)',
            "=INDEX(Table1[Column],1)",
        ):
            with self.subTest(formula=formula):
                self.assertIsNone(recalc.EXTERNAL_REF_RE.search(formula), formula)

    def test_the_reported_location_cap_is_positive(self) -> None:
        self.assertGreater(recalc.MAX_LOCATIONS, 0)


class WorkbookTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self._temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self._temporary.cleanup)
        self.root = Path(self._temporary.name)

    def workbook(self, name: str = "book.xlsx", **cells) -> Path:
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Data"
        for coordinate, value in cells.items():
            sheet[coordinate] = value
        path = self.root / name
        book.save(path)
        book.close()
        return path


class ExternalLinkRiskTests(WorkbookTestCase):
    def test_direct_array_and_sheet_scoped_external_references_are_guarded(self) -> None:
        from openpyxl.workbook.external_link.external import ExternalLink, ExternalBook, ExternalSheetNames
        from openpyxl.packaging.relationship import Relationship
        from openpyxl.workbook.defined_name import DefinedName
        from openpyxl.worksheet.formula import ArrayFormula
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Data"
        link = ExternalLink(externalBook=ExternalBook(
            sheetNames=ExternalSheetNames(sheetName=["Source"]), id="rId1"))
        link.file_link = Relationship(type="externalLinkPath", Target="missing.xlsx", TargetMode="External")
        book._external_links.append(link)
        sheet["A1"] = "='[1]Source'!A1"
        sheet["B1"] = ArrayFormula("B1:B2", "='[1]Source'!A1:A2")
        sheet.defined_names.add(DefinedName("external_input", attr_text="'[1]Source'!$C$1", localSheetId=0))
        sheet["C1"] = "=external_input*2"
        path = self.root / "external.xlsx"
        book.save(path)
        book.close()
        self.assertEqual(recalc.external_links_at_risk(path), ["Data!A1", "Data!B1", "Data!C1"])
        before = path.read_bytes()
        with patch.object(recalc, "_recalc_with_profile") as run:
            result = recalc.recalc(path)
        run.assert_not_called()
        self.assertIn("external_link_cells", result)
        self.assertEqual(path.read_bytes(), before)

    def test_a_workbook_with_no_external_links_is_never_at_risk(self) -> None:
        path = self.workbook(A1=1, A2=2, A3="=SUM(A1:A2)")
        self.assertEqual(recalc.external_links_at_risk(path), [])

    def test_a_non_workbook_file_is_reported_as_no_risk_rather_than_raising(self) -> None:
        path = self.root / "not-a-workbook.xlsx"
        path.write_bytes(b"definitely not a zip")
        self.assertEqual(recalc.external_links_at_risk(path), [])

    def test_a_missing_file_is_reported_as_no_risk(self) -> None:
        self.assertEqual(
            recalc.external_links_at_risk(self.root / "absent.xlsx"), []
        )

    def test_the_check_reads_the_package_without_modifying_it(self) -> None:
        path = self.workbook(A1=1, A2="=A1*2")
        before = path.read_bytes()
        recalc.external_links_at_risk(path)
        self.assertEqual(path.read_bytes(), before)

    def test_a_workbook_without_the_external_links_part_short_circuits(self) -> None:
        # The function returns early unless xl/externalLinks/ exists, so the
        # common case never pays for two full workbook loads.
        path = self.workbook(A1=1)
        with zipfile.ZipFile(path) as archive:
            self.assertFalse(
                any(name.startswith("xl/externalLinks/") for name in archive.namelist())
            )
        self.assertEqual(recalc.external_links_at_risk(path), [])


class MacroSetupTests(WorkbookTestCase):
    def setUp(self) -> None:
        super().setUp()
        import shutil
        if not shutil.which("soffice"):
            self.skipTest("LibreOffice is not installed")

    def test_the_macro_module_is_written_into_the_profile(self) -> None:
        profile = self.root / "profile"
        recalc.setup_libreoffice_macro(profile)
        module = profile / "user" / "basic" / "Standard" / recalc.MACRO_FILENAME
        self.assertTrue(module.is_file(), "macro module was not written")
        self.assertIn("Recalculate", module.read_text(encoding="utf-8"))

    def test_document_url_is_xml_escaped(self) -> None:
        import defusedxml.ElementTree as ET
        profile = self.root / "escaped-profile"
        _, error = recalc.setup_libreoffice_macro(profile, document_url="file:///tmp/a&b.xlsx")
        self.assertIsNone(error)
        macro = ET.parse(profile / "user/basic/Standard" / recalc.MACRO_FILENAME)
        self.assertIn('"file:///tmp/a&b.xlsx"', macro.getroot().text)

    def test_the_macro_xml_is_well_formed(self) -> None:
        import defusedxml.ElementTree as ElementTree

        ElementTree.fromstring(recalc.RECALCULATE_MACRO)

    def test_setup_is_idempotent(self) -> None:
        profile = self.root / "profile"
        recalc.setup_libreoffice_macro(profile)
        first = (profile / "user" / "basic" / "Standard" / recalc.MACRO_FILENAME).read_bytes()
        recalc.setup_libreoffice_macro(profile)
        second = (profile / "user" / "basic" / "Standard" / recalc.MACRO_FILENAME).read_bytes()
        self.assertEqual(first, second)


class RecalcGuardTests(WorkbookTestCase):
    def test_a_missing_soffice_is_reported_clearly(self) -> None:
        import shutil

        if shutil.which("soffice"):
            self.skipTest("LibreOffice is installed; this asserts the absent path")
        path = self.workbook(A1=1, A2="=A1*2")
        result = recalc.recalc(path)
        self.assertIn("soffice", str(result).lower())

    def test_the_missing_soffice_message_names_libreoffice(self) -> None:
        self.assertIn("LibreOffice", recalc.SOFFICE_MISSING)

    def test_the_timestamp_helper_reports_a_files_mtime(self) -> None:
        path = self.workbook(A1=1)
        self.assertEqual(recalc._stamp(path), recalc._stamp(path))
        self.assertIsNotNone(recalc._stamp(path))

    def test_the_timeout_helper_answers_without_raising(self) -> None:
        self.assertIn(recalc.has_gtimeout(), (True, False))

    def test_failed_conversion_does_not_replace_input(self) -> None:
        path = self.workbook(A1=4, A2="=A1*2")
        before = path.read_bytes()
        def fail(staged, *args):
            self.assertEqual(Path(staged).parent, path.parent)
            Path(staged).write_bytes(b"incomplete conversion")
            return {"error": "LibreOffice timed out"}
        with patch.object(recalc, "_recalc_with_profile", side_effect=fail):
            result = recalc.recalc(path)
        self.assertIn("error", result)
        self.assertEqual(path.read_bytes(), before)

    def test_nonpositive_timeout_is_rejected_before_launch(self) -> None:
        path = self.workbook(A1=1)
        with patch.object(recalc, "_recalc_with_profile") as run:
            self.assertIn("error", recalc.recalc(path, timeout=0))
        run.assert_not_called()

    def test_formula_removal_during_conversion_preserves_original(self) -> None:
        import subprocess
        path = self.workbook(A1=3, A2="=A1*2")
        before = path.read_bytes()
        def remove_formula(*args, **kwargs):
            staged = next(self.root.glob("recalc-work-*.xlsx"))
            book = openpyxl.load_workbook(staged)
            book["Data"]["A2"] = 6
            book.save(staged)
            book.close()
            return subprocess.CompletedProcess(args[0], 0, "", "")
        with patch.object(recalc, "setup_libreoffice_macro", return_value=("file:///profile", None)), \
                patch.object(recalc.platform, "system", return_value="Test"), \
                patch.object(recalc.subprocess, "run", side_effect=remove_formula):
            result = recalc.recalc(path)
        self.assertEqual(result["missing_formulas"], ["Data!A2"])
        self.assertEqual(path.read_bytes(), before)


class ResultInspectionTests(WorkbookTestCase):
    def test_error_literals_are_distinct_from_error_cells(self) -> None:
        path = self.workbook(A1="#REF! is a diagnostic", A2="#DIV/0!")
        result = recalc.inspect_results(path)
        self.assertEqual(result["total_errors"], 1)
        self.assertEqual(result["error_summary"]["#DIV/0!"]["locations"], ["Data!A2"])

    def test_missing_formula_cache_is_not_success(self) -> None:
        path = self.workbook(A1=3, A2="=A1*2")
        result = recalc.inspect_results(path)
        self.assertEqual(result["missing_cached_values"], ["Data!A2"])
        self.assertIn("error", result)

    def test_modern_typed_error_is_reported(self) -> None:
        path = self.workbook(A1="placeholder")
        wb = openpyxl.load_workbook(path)
        wb["Data"]["A1"] = "#SPILL!"
        wb["Data"]["A1"].data_type = "e"
        wb.save(path)
        wb.close()
        result = recalc.inspect_results(path)
        self.assertEqual(result["error_summary"]["#SPILL!"]["count"], 1)

    def test_array_formula_anchor_is_counted(self) -> None:
        from openpyxl.worksheet.formula import ArrayFormula
        path = self.workbook(A1=ArrayFormula("A1:A3", "=_xlfn.SEQUENCE(3,1)"))
        self.assertEqual(recalc.formula_locations(path), {("Data", "A1")})

    def test_locations_are_capped_but_count_is_complete(self) -> None:
        path = self.workbook(**{f"A{i}": "#N/A" for i in range(1, recalc.MAX_LOCATIONS + 3)})
        result = recalc.inspect_results(path)
        entry = result["error_summary"]["#N/A"]
        self.assertEqual(entry["count"], recalc.MAX_LOCATIONS + 2)
        self.assertEqual(len(entry["locations"]), recalc.MAX_LOCATIONS)
        self.assertEqual(entry["locations_truncated"], 2)


class LibreOfficeIntegrationTests(WorkbookTestCase):
    def test_recalculate_numeric_empty_string_array_and_error_outputs(self) -> None:
        import shutil
        if not shutil.which("soffice"):
            self.skipTest("LibreOffice is not installed")
        from openpyxl.worksheet.formula import ArrayFormula
        path = self.workbook(A1=4, A2=6, B1="=SUM(A1:A2)",
                             B2='=IF(A1=4,"",0)', B3='="#REF! is text"',
                             B4="=1/0", D1=ArrayFormula("D1:D2", "=A1:A2*2"))
        result = recalc.recalc(path, timeout=60)
        self.assertEqual(result["status"], "errors_found", result)
        self.assertEqual(result["total_errors"], 1)
        self.assertEqual(result["total_formulas"], 5)
        wb = openpyxl.load_workbook(path, data_only=True)
        self.addCleanup(wb.close)
        self.assertEqual(wb["Data"]["B1"].value, 10)
        self.assertIsNone(wb["Data"]["B2"].value)
        self.assertEqual(wb["Data"]["B3"].value, "#REF! is text")
        self.assertEqual([wb["Data"][f"D{i}"].value for i in [1, 2]], [8, 12])


if __name__ == "__main__":
    unittest.main()
