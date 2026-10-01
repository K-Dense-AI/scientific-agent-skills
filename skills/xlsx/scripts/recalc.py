"""
Excel Formula Recalculation Script
Recalculates all formulas in an Excel file using LibreOffice
"""

import argparse
import contextlib
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape as xml_escape

from office.soffice import get_soffice_env, run_soffice

from openpyxl import load_workbook

MACRO_FILENAME = "Module1.xba"
SOFFICE_MISSING = "soffice not found on PATH; LibreOffice is required to recalculate"

MAX_LOCATIONS = 100

EXTERNAL_REF_RE = re.compile(r"""(?<![\w"\[])'?\[\d+\][^!"\[\]]*'?!""")

RECALCULATE_MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      Dim document As Object
      Dim loadArgs(3) As New com.sun.star.beans.PropertyValue
      loadArgs(0).Name = "Hidden"
      loadArgs(0).Value = True
      loadArgs(1).Name = "MacroExecutionMode"
      loadArgs(1).Value = 0
      loadArgs(2).Name = "AsTemplate"
      loadArgs(2).Value = False
      loadArgs(3).Name = "UpdateDocMode"
      loadArgs(3).Value = 0
      document = StarDesktop.loadComponentFromURL("__DOCUMENT_URL__", "_blank", 0, loadArgs())
      document.calculateAll()
      document.store()
      document.close(True)
      StarDesktop.terminate()
    End Sub
</script:module>"""


def has_gtimeout():
    try:
        subprocess.run(
            ["gtimeout", "--version"], capture_output=True, timeout=1, check=False
        )
        return True
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return False


def _stamp(path):
    st = os.stat(path)
    return st.st_mtime_ns, st.st_size


def setup_libreoffice_macro(profile_dir: Path, timeout=30, document_url=""):
    url = profile_dir.as_uri()
    try:
        result = run_soffice(
            ["--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
            capture_output=True,
            text=True,
            timeout=timeout,
        )
    except FileNotFoundError:
        return None, SOFFICE_MISSING
    except subprocess.TimeoutExpired:
        return None, "LibreOffice timed out creating its profile; formulas were NOT recalculated"

    if result.returncode:
        return None, f"LibreOffice profile initialization failed: {result.stderr.strip()}"

    macro_dir = profile_dir / "user" / "basic" / "Standard"
    if not macro_dir.exists():
        return None, "LibreOffice did not create a usable profile; formulas were NOT recalculated"

    try:
        macro = RECALCULATE_MACRO.replace("__DOCUMENT_URL__", xml_escape(document_url))
        (macro_dir / MACRO_FILENAME).write_text(macro, encoding="utf-8")
    except OSError as e:
        return None, f"Could not install the recalculation macro: {e}"

    return url, None


def external_links_at_risk(filename):
    try:
        with zipfile.ZipFile(filename) as archive:
            names = archive.namelist()
    except (zipfile.BadZipFile, OSError):
        return []
    if not any(n.startswith("xl/externalLinks/") for n in names):
        return []

    with contextlib.ExitStack() as stack:
        formulas = load_workbook(filename, data_only=False)
        stack.callback(formulas.close)
        values = load_workbook(filename, data_only=True)
        stack.callback(values.close)

        external_names = [
            name
            for name, dn in formulas.defined_names.items()
            if isinstance(getattr(dn, "value", None), str) and EXTERNAL_REF_RE.search(dn.value)
        ]
        name_re = (
            re.compile(r"\b(" + "|".join(re.escape(n) for n in external_names) + r")\b")
            if external_names
            else None
        )

        at_risk = []
        for sheet in formulas.sheetnames:
            ws = formulas[sheet]
            if not hasattr(ws, "iter_rows"):  
                continue
            cached = values[sheet]
            local_names = [name for name, dn in ws.defined_names.items()
                           if isinstance(getattr(dn, "value", None), str)
                           and EXTERNAL_REF_RE.search(dn.value)]
            local_name_re = (re.compile(r"\b(" + "|".join(re.escape(n) for n in local_names) + r")\b")
                             if local_names else None)
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if cell.data_type == "f" and not isinstance(v, str):
                        v = getattr(v, "text", None)
                    if not (isinstance(v, str) and v.startswith("=")):
                        continue
                    reaches_out = (EXTERNAL_REF_RE.search(v) or (name_re and name_re.search(v))
                                   or (local_name_re and local_name_re.search(v)))
                    if reaches_out and cached[cell.coordinate].value is None:
                        at_risk.append(f"{sheet}!{cell.coordinate}")
        return at_risk


def recalc(filename, timeout=30, force=False):
    if not Path(filename).is_file():
        return {"error": f"File {filename} does not exist"}
    if timeout <= 0:
        return {"error": "Timeout must be positive"}
    if Path(filename).suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return {"error": "Expected an OOXML workbook (.xlsx/.xlsm/.xltx/.xltm)"}

    abs_path = str(Path(filename).resolve())

    if not os.access(abs_path, os.W_OK):
        return {"error": f"{filename} is not writable; recalculation rewrites the file in place"}

    try:
        get_soffice_env()
    except Exception as e:  
        return {"error": f"Could not prepare the LibreOffice environment: {e}"}

    if not force:
        try:
            at_risk = external_links_at_risk(filename)
        except Exception as e:  
            return {"error": f"Could not inspect {filename} for external links: {e}"}
        if at_risk:
            shown = at_risk[:MAX_LOCATIONS]
            return {
                "error": (
                    "Refusing to recalculate: this workbook links to another workbook, and "
                    f"{len(at_risk)} linked cell(s) have lost their cached value (openpyxl strips "
                    "these on save). LibreOffice may replace them with errors or alter external "
                    "links. Restore the source workbook/caches or use a compatible Excel workflow. "
                    "Pass --force only on a disposable copy to accept the risk. Charts and conditional "
                    "formats can hold external references too, so this list may not be exhaustive."
                ),
                "external_link_cells": shown,
                "external_link_cells_truncated": max(0, len(at_risk) - len(shown)),
            }

    # A failed conversion must never partially overwrite the caller's workbook.
    # Stage in the same directory so relative workbook links keep their base
    # and publication is one atomic replace on the same filesystem.
    staged = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix="recalc-work-", suffix=Path(abs_path).suffix,
            dir=Path(abs_path).parent, delete=False,
        ) as output:
            staged = Path(output.name)
        shutil.copy2(abs_path, staged)
        with tempfile.TemporaryDirectory(
            prefix="recalc-lo-profile-", ignore_cleanup_errors=True
        ) as profile_dir:
            result = _recalc_with_profile(staged, str(staged), timeout, Path(profile_dir))
        if "error" not in result:
            os.replace(staged, abs_path)
        return result
    except Exception as e:
        return {"error": f"Recalculation failed; original was not replaced: {e}"}
    finally:
        if staged is not None:
            with contextlib.suppress(OSError):
                staged.unlink(missing_ok=True)
                # A crashed office process can leave this disposable file lock.
                staged.with_name(f".~lock.{staged.name}#").unlink(missing_ok=True)


def _recalc_with_profile(filename, abs_path, timeout, profile_dir: Path):
    started = time.monotonic()
    before_formulas = formula_locations(filename)
    profile_url, err = setup_libreoffice_macro(
        profile_dir, timeout=timeout, document_url=Path(abs_path).as_uri()
    )
    if err:
        return {"error": err}

    timeout = max(5, int(timeout - (time.monotonic() - started)))

    before = _stamp(abs_path)

    cmd = [
        "soffice",
        "--headless",
        "--norestore",
        f"-env:UserInstallation={profile_url}",
        "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
    ]

    if platform.system() == "Linux" and shutil.which("timeout"):
        cmd = ["timeout", str(timeout)] + cmd
    elif platform.system() == "Darwin" and has_gtimeout():
        cmd = ["gtimeout", str(timeout)] + cmd

    timed_out = f"LibreOffice timed out after {timeout}s; original was not replaced. Re-run with a longer timeout."

    try:
        result = subprocess.run(
            cmd, capture_output=True, text=True, env=get_soffice_env(), timeout=timeout + 15
        )
    except subprocess.TimeoutExpired:
        return {"error": timed_out}
    except FileNotFoundError:
        return {"error": SOFFICE_MISSING}

    if result.returncode == 124:
        return {"error": timed_out}

    if result.returncode != 0:
        detail = (result.stderr or "").strip() or f"soffice exited {result.returncode}"
        return {"error": f"LibreOffice failed to recalculate: {detail}"}

    if _stamp(abs_path) == before:
        return {
            "error": (
                "LibreOffice exited cleanly but never rewrote the file, so nothing was "
                "recalculated. Check that no other LibreOffice instance is running, then retry."
            )
        }

    after_formulas = formula_locations(filename)
    missing = before_formulas - after_formulas
    if missing:
        return {"error": "LibreOffice removed formula cells; original was not replaced",
                "missing_formulas": [f"{s}!{c}" for s, c in sorted(missing)[:MAX_LOCATIONS]]}
    return inspect_results(filename, after_formulas)


def formula_locations(filename):
    """Include ordinary formulas and array/data-table anchors (not spill cells)."""
    wb = load_workbook(filename, data_only=False)
    try:
        return {(ws.title, cell.coordinate) for ws in wb.worksheets
                for row in ws.iter_rows() for cell in row if cell.data_type == "f"}
    finally:
        wb.close()


def inspect_results(filename, formulas=None):
    """Read typed errors and cached results without saving through openpyxl."""
    if formulas is None:
        formulas = formula_locations(filename)
    wb = load_workbook(filename, data_only=True)
    error_details = {}
    missing = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    location = f"{ws.title}!{cell.coordinate}"
                    if cell.data_type == "e":
                        error_details.setdefault(str(cell.value), []).append(location)
                    elif ((ws.title, cell.coordinate) in formulas and cell.value is None
                          and cell.data_type not in {"s", "str", "inlineStr"}):
                        # Cached empty strings are legitimate; absent numeric caches are not.
                        missing.append(location)
    finally:
        wb.close()
    if missing:
        return {"error": "Formula cells lack cached results; original was not replaced",
                "missing_cached_values": missing[:MAX_LOCATIONS],
                "missing_cached_values_total": len(missing)}
    total_errors = sum(len(locations) for locations in error_details.values())
    result = {"status": "success" if total_errors == 0 else "errors_found",
              "total_errors": total_errors, "total_formulas": len(formulas), "error_summary": {}}
    for error, locations in error_details.items():
        entry = {"count": len(locations), "locations": locations[:MAX_LOCATIONS]}
        if len(locations) > MAX_LOCATIONS:
            entry["locations_truncated"] = len(locations) - MAX_LOCATIONS
        result["error_summary"][error] = entry
    return result


def main():
    parser = argparse.ArgumentParser(
        description="Recalculate a workbook through LibreOffice; replace only verified output.",
        epilog="JSON status is success or errors_found (both exit 0). An error key exits 1 and preserves the original.",
    )
    parser.add_argument("excel_file")
    parser.add_argument("timeout_seconds", type=int, nargs="?", default=30)
    parser.add_argument("--force", action="store_true", help="Accept risk from external links without cached results")
    args = parser.parse_args()
    result = recalc(args.excel_file, args.timeout_seconds, force=args.force)
    print(json.dumps(result, indent=2))
    sys.exit(1 if "error" in result else 0)


if __name__ == "__main__":
    main()
